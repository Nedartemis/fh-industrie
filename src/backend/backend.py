import json
import os
import re
import shutil
import sys
from pathlib import Path
from typing import Dict, Optional

from docx import Document

from backend.claude_client import ClaudeClient
from backend.excel_manager import ExcelManager
from backend.ocr import ocr_pdf
from frontend.logger import LogLevel
from frontend.logger import log as front_log
from vars import PATH_CONFIG_FILE, PATH_DOCS, PATH_TMP


class Backend:

    template_path: Optional[Path]
    filled_path_file: Optional[Path]

    def __init__(self):
        self.claude_client = ClaudeClient(
            "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"
        )
        # self.filled_path_file = PATH_TMP / "fichier_configuration_rempli.xlsx"
        self.reset()

    def log(self, message: str) -> None:
        print(f"Backend : {message}")

    def reset(self) -> None:
        self.log("Reset")

        # if PATH_CONFIG_FILE.exists():
        #     os.remove(PATH_CONFIG_FILE)

    def extract_infos(self):
        if not PATH_CONFIG_FILE.exists():
            raise RuntimeError("Configuration file does not exist")

        self.log("Extracting infos...")

        # gather files to analyse and info to extract by reading config file
        files_path: Dict[str, str] = {}

        em = ExcelManager(PATH_CONFIG_FILE)
        ws = em.wb.worksheets[1]  # sources
        x, y = 2, 3
        while True:
            name_source = ws.cell(y, x).value
            path_source = ws.cell(y, x + 1).value
            if not name_source or not path_source:
                break
            y += 1

            files_path[name_source] = path_source

        files_infos: Dict[str, List[Tuple[str, str]]] = {}
        ws = em.wb.worksheets[0]  # infos
        x, y = 2, 3
        while True:
            name_info = ws.cell(y, x).value
            description_info = ws.cell(y, x + 1).value
            name_source = ws.cell(y, x + 2).value
            information = ws.cell(y, x + 3).value
            if not name_info:
                break
            y += 1

            if information:
                continue

            if name_source not in files_infos:
                files_infos[name_source] = []
            files_infos[name_source].append((name_info, description_info))

        # error handling (files not supported, ...)
        save = files_path
        files_path = {}
        for name_source, path in save.items():
            if not path.endswith(".pdf"):
                front_log(
                    LogLevel.ERROR,
                    f"L'extension '{Path(path).suffix}' n'est pas supporté en extraction d'information",
                )
            elif not (PATH_TMP / path).exists():
                front_log(
                    LogLevel.ERROR,
                    f"Le fichier {path} n'existe pas dans l'arbre de fichier donné.",
                )
            elif not name_source in files_infos:
                front_log(
                    LogLevel.WARNING,
                    f"Aucune information est à extraire de la source de fichier '{name_source}'.",
                )
            else:
                files_path[name_source] = path

        save = files_infos
        files_infos = {}
        for name_source, e in save.items():
            if name_source not in files_path:
                front_log(
                    LogLevel.ERROR,
                    f"La source '{name_source}' n'a pas de chemin d'accès valide dans l'arborescence.",
                )
                continue
            files_infos[name_source] = e

        # log info
        self.log("Loggin")
        front_log(
            LogLevel.INFO,
            f"{len(files_path)} fichier(s) va/vont être analysé(s) et {sum(len(e) for e in files_infos.values())} information(s) doit/doivent être extraite(s).",
        )

        # extract informations
        all_infos: Dict[str, str] = {}

        for source_name, infos in files_infos.items():
            path_cache = PATH_TMP / (source_name + ".json")
            if path_cache.exists():
                with open(str(path_cache), mode="r") as f:
                    pages = json.load(f)
                front_log(LogLevel.INFO, f"'{source_name}' chargé depuis le cache.")
            else:
                print(files_path[name_source])
                pages = ocr_pdf(PATH_TMP / files_path[name_source])
                front_log(LogLevel.INFO, f"'{source_name}' lu avec un OCR.")
                with open(str(path_cache), mode="w") as f:
                    json.dump(pages, f)

            format_infos = (
                "```json\n{ \n"
                + "\n\t".join([f'"{info}": "..."' for info, _ in infos])
                + "\n}```"
            )
            prompt_system = f"""
                Extrait toutes les informations que tu trouves sous format json :
                {format_infos}

                Si tu ne trouves pas l'info, remplie le champs par "None".
            """

            messages = [{"role": "user", "content": "\n\n".join(pages[:1])}]

            response = self.claude_client.create_message(
                system=prompt_system,
                messages=messages,
                max_tokens=1000,
                temperature=1,
            )

            text_infos = response["content"][0]["text"]
            print(f"text_infos : {text_infos}")
            res = re.search(
                pattern="```json(.*)```", string=text_infos, flags=re.DOTALL
            )
            extracted_str = res.group(1)
            print(f"extracted_str : {extracted_str}")
            all_infos.update(
                {
                    key: value
                    for key, value in json.loads(extracted_str).items()
                    if value != "None"
                }
            )

            for name_info in all_infos:
                if name_info not in [e for e, _ in infos]:
                    front_log(
                        LogLevel.WARNING,
                        f"'{name_info}' a été extrait mais n'avait pas été demandé.",
                    )

        front_log(
            LogLevel.INFO,
            f"{len(all_infos)} informations ont été extraites avec succès.",
        )

        # copy config file
        self.filled_path_file = PATH_TMP / "fichier_configuration_rempli.xlsx"

        # fill config file
        em = ExcelManager(PATH_CONFIG_FILE)
        ws = em.wb.worksheets[0]  # infos
        x, y = 2, 3
        n = 0
        while True:
            name_info = ws.cell(y, x).value
            if not name_info:
                break
            y += 1
            if not name_info in all_infos:
                continue

            n += 1
            ws.cell(y - 1, x + 3, value=all_infos[name_info])
        em.save(self.filled_path_file)
        em.wb.close()

        return str(self.filled_path_file)

    def set_template_path(self, path: str):
        self.template_path = Path(path)
        print(self.template_path)

    def fill_template(self):

        if not self.template_path or not self.template_path.exists():
            raise RuntimeError("Template file does not exist")

        if not self.filled_path_file or not self.filled_path_file.exists():
            raise RuntimeError("Filled file does not exist")

        self.log("Filling template...")

        # extract infos from filled file
        em = ExcelManager(self.filled_path_file)

        infos: Dict[str, str] = {}
        ws = em.wb.worksheets[0]  # infos
        x, y = 2, 3
        while True:
            name_info = ws.cell(y, x).value
            information = ws.cell(y, x + 3).value
            if not name_info:
                break

            y += 1
            if not information:
                continue

            infos[name_info] = information

        print(infos)

        # copy and filled the template
        path_output = PATH_TMP / (
            self.template_path.stem + "_rempli" + self.template_path.suffix
        )

        if self.template_path.suffix.endswith("xlsx"):
            front_log(LogLevel.INFO, "Le modèle est un excel")

            excel = ExcelManager(self.template_path)
            excel.replace_content(infos)
            excel.save(path_output)
        elif self.template_path.suffix.endswith("docx"):
            front_log(LogLevel.INFO, "Le modèle est un .docx")
            doc = Document(self.template_path)
            for para in doc.paragraphs:
                for key, value in infos.items():
                    if key in para.text:
                        para.text = para.text.replace("{" + key + "}", value)
                        print(para.text)

            doc.save(path_output)
        else:
            front_log(
                LogLevel.WARNING,
                f"L'extension '{self.template_path.suffix}' du modèle n'est pas supporté.",
            )
            return None

        front_log(LogLevel.INFO, "Le modèle a bien été rempli et sauvegardé.")
        return path_output


def _main():
    Backend().extract_infos()


if __name__ == "__main__":
    _main()
